#!/usr/bin/env python3
"""Build the NYC job-search workbook and CSV from data/jobs.jsonl and data/excluded.jsonl.

Each line of data/jobs.jsonl is one posting. Short keys keep the data file compact:
  co, title, cat, ind, loc, mode, base, ote, var, b_lo, b_hi, o_lo, o_hi (in $K),
  src_ev, yrs, ng, tr, why, fit, hire, grow, apply, src, status, ver, dl, flags,
  next, stretch, ecom, vl, revisit
"""
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
OUT_XLSX = ROOT / "NYC_Sales_GTM_Jobs_Master.xlsx"
OUT_CSV = ROOT / "NYC_Sales_GTM_Jobs_All.csv"

CATEGORIES = {
    "A": "AI / B2B SaaS sales",
    "B": "Startup GTM / commercial",
    "C": "E-commerce & merchant sales",
    "D": "AdTech / MarTech / creator commerce",
    "E": "Consumer, fashion & marketplaces",
    "F": "Other high-upside commercial",
}


def norm_company(name):
    s = name.lower().replace("&", "and")
    s = re.sub(r"\b(inc|llc|ltd|corp|corporation|co|company|technologies|labs|hq)\b\.?", "", s)
    return re.sub(r"[^a-z0-9]", "", s)


def norm_title(title):
    s = title.lower()
    s = re.sub(r"\(.*?\)", "", s)
    s = re.sub(r"[-–,|]\s*(new york|nyc|ny|hybrid|remote|in[- ]office).*$", "", s)
    s = s.replace("&", "and")
    return re.sub(r"[^a-z0-9]", "", s)


def load_jsonl(path):
    rows = []
    if not path.exists():
        return rows
    for i, line in enumerate(path.read_text().splitlines(), 1):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        try:
            rows.append(json.loads(line))
        except json.JSONDecodeError as e:
            sys.exit(f"{path.name}:{i}: bad JSON: {e}")
    return rows


def mid(lo, hi):
    vals = [v for v in (lo, hi) if isinstance(v, (int, float))]
    return sum(vals) / len(vals) if vals else None


def tier(r):
    o_mid, b_mid = mid(r.get("o_lo"), r.get("o_hi")), mid(r.get("b_lo"), r.get("b_hi"))
    top = max([v for v in (r.get("o_hi"), r.get("o_lo"), r.get("b_hi"), r.get("b_lo")) if isinstance(v, (int, float))], default=None)
    if o_mid is not None:
        if o_mid >= 120:
            return "A ($120K+ OTE)"
        if o_mid >= 100:
            return "B ($100K-$119K OTE)"
        if o_mid >= 90:
            return "C ($90K-$99K OTE)"
    elif b_mid is not None and b_mid >= 90:
        return "D (Salary $90K+, no OTE disclosed)"
    if top is None:
        return "Unverified"
    if top >= 90:
        return "Near target (range reaches $90K+, midpoint below)"
    return "Below target"


def comp_score(r):
    o_mid, b_mid = mid(r.get("o_lo"), r.get("o_hi")), mid(r.get("b_lo"), r.get("b_hi"))
    if o_mid is None and b_mid is None:
        return 4.0
    v = o_mid if o_mid is not None else b_mid * 1.1  # salary-only roles: all guaranteed
    for floor, score in ((150, 10), (135, 9.3), (120, 8.5), (110, 7.5), (100, 6.5), (90, 5.5), (80, 4), (0, 2)):
        if v >= floor:
            s = score
            break
    b_lo = r.get("b_lo")
    if isinstance(b_lo, (int, float)) and o_mid is not None and b_lo < 55:
        s -= 0.5  # low guaranteed base behind a big OTE
    return round(max(s, 0), 1)


def money(lo, hi):
    def f(v):
        return f"${v:,.0f}K" if isinstance(v, (int, float)) else ""
    if lo is None and hi is None:
        return ""
    if lo is None or hi is None or lo == hi:
        return f(lo if lo is not None else hi)
    return f"{f(lo)}-{f(hi)}"


COLUMNS = [
    ("Unique ID", "id", 9),
    ("Priority Rank", "rank", 8),
    ("Company", "co", 22),
    ("Exact Job Title", "title", 34),
    ("Job Category", "cat_label", 22),
    ("Industry", "ind", 22),
    ("NYC Office Location", "loc", 24),
    ("Work Arrangement", "mode", 14),
    ("Guaranteed Base Salary", "base_txt", 18),
    ("Advertised OTE", "ote_txt", 18),
    ("Commission / Variable Pay Details", "var", 30),
    ("Compensation Tier", "tier", 20),
    ("Base Midpoint ($K, sortable)", "b_mid", 11),
    ("OTE Midpoint ($K, sortable)", "o_mid", 11),
    ("Salary Source & Evidence", "src_ev", 34),
    ("Required Years of Experience", "yrs", 14),
    ("New Grad / 2027 Eligibility & Start Date", "ng", 28),
    ("Transferable Experience Accepted?", "tr", 22),
    ("Why It Fits Me", "why", 40),
    ("Fit Score (/10)", "fit", 7),
    ("Hireability Score (/10)", "hire", 7),
    ("Comp Attractiveness (/10)", "comp", 7),
    ("Career Growth (/10)", "grow", 7),
    ("Overall Priority (/100)", "overall", 8),
    ("Stretch Role?", "stretch_txt", 8),
    ("Direct Application URL", "apply", 40),
    ("Source URL", "src", 40),
    ("Posting Status", "status", 22),
    ("Verification Method", "vl", 26),
    ("Date/Time Verified (ET)", "ver", 17),
    ("Application Deadline", "dl", 14),
    ("Red Flags / Caveats", "flags", 40),
    ("Recommended Next Action", "next", 34),
]


def enrich(r):
    r = dict(r)
    for k in ("fit", "hire", "grow"):
        r[k] = float(r.get(k) or 0)
    r["comp"] = comp_score(r)
    r["overall"] = round(r["fit"] * 3 + r["hire"] * 3 + r["comp"] * 2 + r["grow"] * 2, 1)
    r["tier"] = tier(r)
    r["b_mid"] = mid(r.get("b_lo"), r.get("b_hi"))
    r["o_mid"] = mid(r.get("o_lo"), r.get("o_hi"))
    r["base_txt"] = r.get("base") or money(r.get("b_lo"), r.get("b_hi")) or "Not disclosed"
    r["ote_txt"] = r.get("ote") or money(r.get("o_lo"), r.get("o_hi")) or "Not disclosed"
    r["cat_label"] = f'{r.get("cat", "?")}: {CATEGORIES.get(r.get("cat"), "Other")}'
    r["stretch_txt"] = "STRETCH" if r.get("stretch") else ""
    return r


def main():
    suppress = {norm_company(l) for l in (ROOT / "suppression.txt").read_text().splitlines() if l.strip()}
    raw = load_jsonl(DATA / "jobs.jsonl")
    excluded = [dict(x, reason=x.get("reason", "")) for x in load_jsonl(DATA / "excluded.jsonl")]

    active, seen = [], {}
    for r in raw:
        key = (norm_company(r["co"]), norm_title(r["title"]))
        if key in seen:
            excluded.append(dict(r, reason=f"Duplicate of {seen[key]} (same company + title)"))
            continue
        if norm_company(r["co"]) in suppress and not r.get("revisit"):
            excluded.append(dict(r, reason="On suppression list (already researched)"))
            continue
        r = enrich(r)
        if r["tier"] == "Below target":
            excluded.append(dict(r, reason="Advertised pay below $90K target at every point of the range"))
            continue
        seen[key] = r["co"] + " / " + r["title"]
        active.append(r)

    active.sort(key=lambda r: -r["overall"])
    for i, r in enumerate(active, 1):
        r["rank"] = i
        # stable ID derived from company + title so it survives re-sorting and rebuilds
        digest = hashlib.sha1(f'{norm_company(r["co"])}|{norm_title(r["title"])}'.encode()).hexdigest()
        r["id"] = "J" + digest[:6].upper()

    paid = lambda r: r["tier"][0] in "ABCD"
    tabs = {
        "Priority Applications": [r for r in active if paid(r) and r["overall"] >= 68][:75],
        "All Verified NYC Jobs": active,
        "$120K+ OTE": [r for r in active if r["tier"].startswith("A")],
        "Early-Career Friendly": [r for r in active if r["hire"] >= 7 and not r.get("stretch")],
        "E-Commerce & Consumer Tech": [r for r in active if r.get("ecom")],
        "Unverified Compensation": [r for r in active if r["tier"] == "Unverified"],
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    tiers = Counter(r["tier"] for r in active)
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = [
        ("NYC Sales / AI / GTM / E-Commerce Job Search — Master Workbook", True),
        (f"Last rebuilt: {now} (container clock, UTC)", False),
        ("", False),
        (f"Unique active NYC openings in database: {len(active)}", True),
    ]
    for t in sorted(tiers):
        lines.append((f"   {t}: {tiers[t]}", False))
    lines += [
        (f"Excluded / closed / duplicate / suppressed rows: {len(excluded)}", False),
        ("", False),
        ("How to read this workbook", True),
        ("• Tier uses the MIDPOINT of the advertised range (OTE if disclosed, otherwise base). OTE is never guaranteed; base is shown separately.", False),
        ("• Overall Priority (/100) = Fit×3 + Hireability×3 + Comp×2 + Growth×2. Comp score is computed from the advertised numbers; unverified comp scores 4/10.", False),
        ("• 'Near target' = the advertised range reaches $90K+ but its midpoint is below $90K. Kept for review, never put in Priority.", False),
        ("• Verification: this cloud session could not open job pages directly (network policy). Every row was checked through web-search results that index the posting's own text. "
         "'Official ATS' means the indexed page was the employer's own careers/ATS page; 'Aggregator' means Built In, LinkedIn, Wellfound, etc. Re-open the link before applying.", False),
        ("• STRETCH = requirements exceed a new grad's profile (e.g., 2+ yrs closing). Included only when exceptional.", False),
        ("• Suppression list (companies you already researched) is applied automatically; matches appear in Excluded / Closed.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if i == 1 else 11)
    ws.column_dimensions["A"].width = 150

    head_fill = PatternFill("solid", fgColor="1F3864")
    for name, rows in tabs.items():
        sh = wb.create_sheet(name[:31])
        for j, (label, _, width) in enumerate(COLUMNS, 1):
            c = sh.cell(row=1, column=j, value=label)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
            sh.column_dimensions[get_column_letter(j)].width = width
        for i, r in enumerate(rows, 2):
            for j, (_, key, _) in enumerate(COLUMNS, 1):
                v = r.get(key, "")
                if isinstance(v, bool):
                    v = "Yes" if v else ""
                c = sh.cell(row=i, column=j, value=v if v is not None else "")
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if key in ("apply", "src") and isinstance(v, str) and v.startswith("http"):
                    c.hyperlink = v
                    c.font = Font(color="0563C1", underline="single")
        sh.freeze_panes = "E2"
        sh.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"
        sh.row_dimensions[1].height = 45

    ex_cols = [("Company", "co", 24), ("Job Title", "title", 36), ("Reason Excluded", "reason", 50),
               ("Advertised Pay", "ote", 22), ("Source URL", "src", 50), ("Checked (ET)", "ver", 17)]
    sh = wb.create_sheet("Excluded - Closed")
    for j, (label, _, width) in enumerate(ex_cols, 1):
        c = sh.cell(row=1, column=j, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        sh.column_dimensions[get_column_letter(j)].width = width
    for i, r in enumerate(excluded, 2):
        for j, (_, key, _) in enumerate(ex_cols, 1):
            v = r.get(key) or (money(r.get("o_lo"), r.get("o_hi")) if key == "ote" else "")
            sh.cell(row=i, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    sh.freeze_panes = "A2"
    sh.auto_filter.ref = f"A1:F{max(len(excluded) + 1, 2)}"

    wb.save(OUT_XLSX)
    with OUT_CSV.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow([c[0] for c in COLUMNS])
        for r in active:
            w.writerow(["Yes" if r.get(k) is True else r.get(k, "") for _, k, _ in COLUMNS])

    print(f"active={len(active)} excluded={len(excluded)} tiers={dict(tiers)}")
    print(f"tabs: " + ", ".join(f"{k}={len(v)}" for k, v in tabs.items()))


if __name__ == "__main__":
    main()
